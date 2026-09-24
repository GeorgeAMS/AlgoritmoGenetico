from pathlib import Path

from PIL import Image as PilImage, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Flowable,
    Image,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)


RAIZ = Path(__file__).resolve().parents[1]
CARPETA = Path(__file__).resolve().parent
SALIDA_DOCX = RAIZ / "Documentacion_AlgoritmosGeneticos.docx"
SALIDA_PDF = RAIZ / "Documentacion_AlgoritmosGeneticos.pdf"
CAPTURA_LISTAS = CARPETA / "vista_aplicacion.png"
CAPTURA_NAVE = CARPETA / "vista_aterrizaje.png"
CROMOSOMA_LISTAS = CARPETA / "cromosoma.png"
CROMOSOMA_NAVE = CARPETA / "cromosoma_nave.png"

TITULO = (
    "Algoritmos Genéticos: listas simplemente ligadas "
    "y aterrizaje evolutivo"
)
AUTORES = (
    "Jorge Andrés Martínez Santos  |  Marián Garcés Delgado  |  "
    "Mariana Sinisterra"
)
INSTITUCION = "Politécnico Colombiano Jaime Isaza Cadavid — Medellín, Colombia"
REPOSITORIO = "https://github.com/GeorgeAMS/AlgoritmoGenetico"

DESCRIPCION_GENERAL = [
    (
        "Punto 1 — Listas simplemente ligadas. El problema consiste en encontrar "
        "una cadena de operaciones que transforme "
        "un valor inicial de cero en un número objetivo definido por el usuario. "
        "Cada posición de la cadena puede contener una de tres operaciones: sumar "
        "uno (+1), sumar dos (+2) o multiplicar por dos (×2). Por ejemplo, para "
        "alcanzar el objetivo 11 se puede utilizar la secuencia +1, ×2, ×2, +1, "
        "×2, +1. Aunque una persona puede resolver una instancia pequeña por "
        "prueba y error, el número de cadenas posibles crece exponencialmente al "
        "aumentar la longitud del cromosoma."
    ),
    (
        "La finalidad del algoritmo genético es optimizar la cercanía entre el "
        "resultado producido por cada cadena y el objetivo solicitado. La "
        "particularidad del ejercicio es que cada cromosoma se implementa como una "
        "lista simplemente ligada: cada nodo almacena una operación y una "
        "referencia al siguiente nodo. La población también se almacena mediante "
        "nodos enlazados. La aplicación permite modificar población, generaciones, "
        "cruce, mutación y elitismo, y conserva en una tabla todas las cadenas "
        "evaluadas para observar cómo se construye el mejor resultado."
    ),
    (
        "Punto 2 — Aterrizaje. El segundo problema consiste en evolucionar el "
        "control de una nave que comienza en el aire y debe llegar a una pista. "
        "Cada individuo contiene doce acciones; una acción combina un giro "
        "(-1, 0 o 1) y la decisión de activar o no el propulsor. La simulación "
        "aplica gravedad, empuje, velocidad y ángulo durante ocho instantes por "
        "cada acción."
    ),
    (
        "En este punto se optimizan simultáneamente la cercanía al centro de la "
        "pista, las velocidades horizontal y vertical, y la inclinación de la "
        "nave. No basta con tocar el suelo: para considerarse un aterrizaje, debe "
        "caer sobre la pista, hacerlo lentamente y conservar un ángulo seguro. "
        "La interfaz permite cambiar población, cruce, mutación, elitismo y "
        "velocidad de visualización."
    ),
]

TRABAJOS_RELACIONADOS = [
    (
        "Holland [1] formuló las bases de los algoritmos genéticos y explicó cómo "
        "una población codificada puede evolucionar mediante selección y "
        "recombinación. Su propuesta sustenta el ciclo utilizado en este trabajo: "
        "crear individuos, evaluarlos, seleccionar progenitores y producir nuevas "
        "generaciones hasta cumplir una condición de parada."
    ),
    (
        "Goldberg [2] presentó aplicaciones prácticas de los algoritmos genéticos "
        "en búsqueda y optimización, además de analizar la presión selectiva, el "
        "cruce y la mutación. Estos conceptos se relacionan directamente con la "
        "selección por torneo y con las probabilidades configurables de cruce y "
        "mutación usadas para evolucionar las cadenas de operaciones."
    ),
    (
        "Mitchell [3] destacó que la representación del cromosoma y la función de "
        "aptitud determinan gran parte del comportamiento de un algoritmo genético. "
        "En esta aplicación, la representación no es un arreglo tradicional, sino "
        "una lista simplemente ligada. Goodrich, Tamassia y Goldwasser [4] "
        "describen esta estructura como una secuencia de nodos conectados mediante "
        "una única referencia, principio empleado tanto en los cromosomas como en "
        "la población."
    ),
    (
        "Michalewicz y Fogel [5] estudiaron el uso de metaheurísticas en problemas "
        "donde una solución es una secuencia de decisiones. Ese enfoque se "
        "relaciona con el punto de aterrizaje, porque el cromosoma no representa "
        "una posición final aislada, sino un programa de acciones de giro y empuje "
        "que produce una trayectoria completa."
    ),
]

FUNCION_OBJETIVO = (
    "Punto 1. La función objetivo asigna mayor aptitud a las cadenas cuyo resultado se "
    "encuentra más cerca del valor objetivo. Sea O el objetivo indicado por el "
    "usuario y R(c) el resultado de ejecutar, desde cero, todas las operaciones "
    "del cromosoma c. Se resta la distancia absoluta a una base de 100 y se "
    "establece 1 como valor mínimo para que todos los individuos puedan participar "
    "en la selección. Cuando R(c) es igual a O, el fitness alcanza su máximo de 100."
)

FUNCION_OBJETIVO_NAVE = (
    "Punto 2. La aptitud de la nave combina cuatro términos continuos: distancia "
    "horizontal al centro de la pista, velocidad vertical, velocidad horizontal "
    "y ángulo. Además, suma bonificaciones por tocar el suelo, hacerlo sobre la "
    "pista y cumplir completamente las condiciones de aterrizaje. Por tanto, una "
    "trayectoria puede mejorar gradualmente antes de conseguir un aterrizaje."
)

DEFINICION_POBLACION = [
    (
        "Punto 1. La población es una ListaPoblacion formada por objetos NodoIndividuo. "
        "Cada nodo guarda un Individuo y la referencia al siguiente. Un individuo "
        "contiene su ListaOperaciones, el resultado obtenido y el fitness. Con la "
        "configuración inicial se generan 40 cadenas aleatorias, aunque el usuario "
        "puede elegir entre 2 y 200 individuos desde la interfaz."
    ),
    (
        "El cromosoma es una ListaOperaciones. Cada NodoOperacion contiene uno de "
        "los genes válidos (+1, +2 o ×2) y un enlace siguiente; el último enlace "
        "vale null. La Figura 1 presenta una cadena de seis nodos que transforma "
        "cero en once. Esta representación obliga a recorrer la lista desde la "
        "cabeza para evaluar, cruzar, mutar o mostrar el individuo."
    ),
    (
        "Punto 2. La población inicial contiene 60 naves. Cada cromosoma es una "
        "secuencia de doce genes y cada gen almacena dos valores: rotación "
        "(-1 izquierda, 0 sin giro o 1 derecha) y propulsión (0 apagada o 1 "
        "encendida). La simulación conserva para cada individuo su trayectoria, "
        "posición final, velocidades, ángulo, estado y fitness."
    ),
    (
        "La Figura 2 representa un cromosoma de control. Los genes se ejecutan "
        "de izquierda a derecha y cada uno se mantiene durante ocho pasos físicos. "
        "De este modo, el orden de las acciones modifica completamente la ruta de "
        "la nave y el instante en que utiliza el combustible."
    ),
]

CRUCE_MUTACION = [
    (
        "Punto 1. El cruce utiliza un punto de corte aleatorio situado entre el primer y el "
        "último nodo. El hijo recibe los nodos anteriores al corte del padre y los "
        "nodos restantes de la madre. Las operaciones se copian en nodos nuevos; "
        "de esta manera, padres e hijos no comparten referencias y una modificación "
        "posterior no altera accidentalmente otro cromosoma. El cruce solo se "
        "aplica cuando un número aleatorio cumple la probabilidad configurada."
    ),
    (
        "La mutación selecciona una posición aleatoria y sustituye su operación por "
        "otra diferente entre +1, +2 y ×2. Esta variación permite introducir "
        "información que no estaba presente en los padres y ayuda a escapar de "
        "resultados estancados. Antes de crear hijos se conservan los N mejores "
        "individuos mediante elitismo; los demás progenitores se eligen por torneos "
        "de tres candidatos."
    ),
    (
        "Punto 2. El cruce de las naves también utiliza un punto de corte: copia "
        "las primeras acciones de un progenitor y las restantes del otro. Los "
        "padres se escogen mediante torneo entre individuos de la parte superior "
        "de la población ordenada por fitness. Los mejores individuos indicados "
        "por el elitismo pasan completos a la siguiente generación."
    ),
    (
        "La mutación se evalúa de forma independiente sobre cada gen del control. "
        "Cuando se cumple la probabilidad Pm, se genera una nueva combinación de "
        "giro y propulsión. Esta mutación puntual puede cambiar la orientación o "
        "el empuje en un momento específico sin reemplazar toda la trayectoria."
    ),
]

APLICACION = (
    "Punto 1. La aplicación fue construida en Java con programación orientada a objetos. "
    "Las estructuras NodoOperacion, ListaOperaciones, NodoIndividuo y "
    "ListaPoblacion fueron implementadas manualmente. La interfaz utiliza Java "
    "Swing y ejecuta la evolución en un SwingWorker para evitar que la ventana se "
    "bloquee. El usuario puede cambiar objetivo, nodos, población, generaciones, "
    "cruce, mutación, elitismo y pausa. El panel superior dibuja la mejor lista y "
    "la tabla inferior registra todas las cadenas creadas, su resultado y fitness. "
    "En la Figura 3 se presenta la aplicación durante una ejecución."
)

APLICACION_NAVE = (
    "Punto 2. La aplicación de aterrizaje fue construida con HTML, CSS y "
    "JavaScript. Dos elementos canvas dibujan la simulación y la gráfica de "
    "fitness. El panel lateral contiene controles para población, Pc, Pm, "
    "elitismo y velocidad. Durante la ejecución se presentan la generación, el "
    "estado de la mejor nave, el número de aterrizajes y el mayor fitness. En la "
    "Figura 4 se muestra la interfaz de esta aplicación."
)

CONCLUSIONES = [
    (
        "El algoritmo genético permite encontrar cadenas cercanas o iguales al "
        "objetivo sin enumerar manualmente todas las combinaciones. La calidad de "
        "la búsqueda depende del equilibrio entre conservación y diversidad: un "
        "elitismo alto protege soluciones buenas, mientras que el cruce y la "
        "mutación exploran nuevas cadenas."
    ),
    (
        "La implementación con listas simplemente ligadas demuestra que un "
        "cromosoma no tiene que almacenarse necesariamente en un arreglo. Sin "
        "embargo, acceder a una posición requiere recorrer los enlaces anteriores, "
        "por lo que la estructura hace visibles tanto las ventajas de una "
        "representación dinámica como su coste de recorrido."
    ),
    (
        "Mostrar todas las cadenas de todas las generaciones facilita la "
        "explicación del proceso evolutivo. La aplicación no presenta únicamente "
        "la solución final: permite comparar individuos, identificar élites y "
        "observar el efecto de cambiar población, mutación y cruce."
    ),
    (
        "En el segundo punto, una función de fitness gradual resulta necesaria "
        "porque los aterrizajes válidos son poco frecuentes al inicio. Premiar "
        "la proximidad, la estabilidad y la reducción de velocidad guía a la "
        "población hacia una solución antes de que alguna nave aterrice."
    ),
    (
        "Los dos puntos comparten selección, cruce, mutación y elitismo, pero "
        "demuestran que la representación y el fitness deben diseñarse según el "
        "problema. Una lista de operaciones busca un valor exacto; una secuencia "
        "de control debe equilibrar varias condiciones físicas."
    ),
]

BIBLIOGRAFIA = [
    "[1] J. H. Holland, Adaptation in Natural and Artificial Systems. University of Michigan Press, 1975.",
    "[2] D. E. Goldberg, Genetic Algorithms in Search, Optimization, and Machine Learning. Addison-Wesley, 1989.",
    "[3] M. Mitchell, An Introduction to Genetic Algorithms. MIT Press, 1996.",
    "[4] M. T. Goodrich, R. Tamassia y M. H. Goldwasser, Data Structures and Algorithms in Java, 6.ª ed. Wiley, 2014.",
    "[5] Z. Michalewicz y D. B. Fogel, How to Solve It: Modern Heuristics, 2.ª ed. Springer, 2004.",
]


def fuente_pil(nombre: str, tamanio: int):
    return ImageFont.truetype(f"C:/Windows/Fonts/{nombre}", tamanio)


def generar_figura_cromosoma() -> None:
    imagen = PilImage.new("RGB", (1500, 300), "white")
    dibujo = ImageDraw.Draw(imagen)
    operaciones = ["+1", "×2", "×2", "+1", "×2", "+1"]
    valores = ["0 → 1", "1 → 2", "2 → 4", "4 → 5", "5 → 10", "10 → 11"]
    x, y = 55, 90
    ancho, alto, espacio = 155, 100, 62
    dibujo.text((55, 35), "cabeza", font=fuente_pil("arialbd.ttf", 24), fill="#0369A1")

    for operacion, valor in zip(operaciones, valores):
        dibujo.rounded_rectangle(
            (x, y, x + ancho, y + alto),
            radius=18,
            fill="#E0F2FE",
            outline="#0284C7",
            width=4,
        )
        caja = dibujo.textbbox((0, 0), operacion, font=fuente_pil("arialbd.ttf", 32))
        dibujo.text(
            (x + (ancho - (caja[2] - caja[0])) / 2, y + 13),
            operacion,
            font=fuente_pil("arialbd.ttf", 32),
            fill="#0F172A",
        )
        caja = dibujo.textbbox((0, 0), valor, font=fuente_pil("arial.ttf", 18))
        dibujo.text(
            (x + (ancho - (caja[2] - caja[0])) / 2, y + 62),
            valor,
            font=fuente_pil("arial.ttf", 18),
            fill="#334155",
        )
        inicio = x + ancho
        fin = inicio + espacio - 12
        centro = y + alto // 2
        dibujo.line((inicio + 5, centro, fin, centro), fill="#64748B", width=4)
        dibujo.polygon(
            [(fin, centro), (fin - 13, centro - 8), (fin - 13, centro + 8)],
            fill="#64748B",
        )
        x += ancho + espacio

    dibujo.rounded_rectangle(
        (x - 7, y + 20, x + 92, y + 80), radius=12, fill="#CBD5E1"
    )
    dibujo.text(
        (x + 8, y + 36),
        "null",
        font=fuente_pil("consola.ttf", 22),
        fill="#0F172A",
    )
    dibujo.text(
        (55, 235),
        "Resultado = 11     Fitness = 100",
        font=fuente_pil("arialbd.ttf", 22),
        fill="#047857",
    )
    imagen.save(CROMOSOMA_LISTAS)


def generar_figura_cromosoma_nave() -> None:
    imagen = PilImage.new("RGB", (1500, 360), "white")
    dibujo = ImageDraw.Draw(imagen)
    genes = [
        "(0,1)", "(-1,1)", "(0,0)", "(1,1)", "(0,1)", "(0,0)",
        "(1,0)", "(0,1)", "(-1,0)", "(0,1)", "(0,0)", "(1,1)",
    ]
    ancho, alto, espacio = 180, 82, 35
    for indice, gen in enumerate(genes):
        fila, columna = divmod(indice, 6)
        x = 55 + columna * (ancho + espacio)
        y = 55 + fila * 135
        dibujo.rounded_rectangle(
            (x, y, x + ancho, y + alto),
            radius=16,
            fill="#EDE9FE",
            outline="#7C3AED",
            width=4,
        )
        dibujo.text(
            (x + 20, y + 15),
            f"Gen {indice + 1}",
            font=fuente_pil("arialbd.ttf", 19),
            fill="#4C1D95",
        )
        dibujo.text(
            (x + 85, y + 44),
            gen,
            font=fuente_pil("consola.ttf", 21),
            fill="#0F172A",
        )
        if columna < 5:
            inicio, fin, centro = x + ancho + 4, x + ancho + espacio - 4, y + alto // 2
            dibujo.line((inicio, centro, fin, centro), fill="#64748B", width=3)
            dibujo.polygon(
                [(fin, centro), (fin - 10, centro - 6), (fin - 10, centro + 6)],
                fill="#64748B",
            )
    dibujo.text(
        (55, 320),
        "Cada gen = (rotación, propulsión)",
        font=fuente_pil("arialbd.ttf", 21),
        fill="#5B21B6",
    )
    imagen.save(CROMOSOMA_NAVE)


def configurar_docx(documento: Document) -> None:
    normal = documento.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    for nombre, tamanio in (("Title", 18), ("Heading 1", 14)):
        estilo = documento.styles[nombre]
        estilo.font.name = "Arial"
        estilo.font.size = Pt(tamanio)
        estilo.font.bold = True

    seccion = documento.sections[0]
    seccion.top_margin = Inches(0.75)
    seccion.bottom_margin = Inches(0.75)
    seccion.left_margin = Inches(0.85)
    seccion.right_margin = Inches(0.85)


def texto_docx(documento: Document, contenido: str) -> None:
    parrafo = documento.add_paragraph(contenido)
    parrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def titulo_figura_docx(documento: Document, texto: str) -> None:
    parrafo = documento.add_paragraph(texto)
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    parrafo.runs[0].italic = True


def generar_docx() -> None:
    documento = Document()
    configurar_docx(documento)

    titulo = documento.add_paragraph()
    titulo.style = "Title"
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo.add_run(TITULO)
    autores = documento.add_paragraph(AUTORES)
    autores.alignment = WD_ALIGN_PARAGRAPH.CENTER
    institucion = documento.add_paragraph(INSTITUCION)
    institucion.alignment = WD_ALIGN_PARAGRAPH.CENTER

    documento.add_heading("1. Descripción general", level=1)
    for parrafo in DESCRIPCION_GENERAL:
        texto_docx(documento, parrafo)

    documento.add_heading("2. Trabajos relacionados", level=1)
    for parrafo in TRABAJOS_RELACIONADOS:
        texto_docx(documento, parrafo)

    documento.add_heading("3. Función objetivo", level=1)
    texto_docx(documento, FUNCION_OBJETIVO)
    ecuacion = documento.add_paragraph(
        "f(c) = max(1, 100 − |O − R(c)|)"
    )
    ecuacion.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ecuacion.runs[0].font.name = "Cambria Math"
    ecuacion.runs[0].font.size = Pt(13)
    titulo_figura_docx(documento, "Ecuación 1. Función objetivo del punto 1.")
    texto_docx(documento, FUNCION_OBJETIVO_NAVE)
    ecuacion = documento.add_paragraph(
        "F = 4000/(1 + 0,15|x − c|) + 1500/(1 + |vy|) + "
        "600/(1 + |vx|) + 400/(1 + 2|a|) + B"
    )
    ecuacion.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ecuacion.runs[0].font.name = "Cambria Math"
    ecuacion.runs[0].font.size = Pt(11)
    titulo_figura_docx(documento, "Ecuación 2. Función objetivo del punto 2.")

    documento.add_heading("4. Definición de población", level=1)
    for parrafo in DEFINICION_POBLACION[:2]:
        texto_docx(documento, parrafo)
    documento.add_picture(str(CROMOSOMA_LISTAS), width=Inches(6.65))
    documento.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo_figura_docx(documento, "Figura 1. Cromosoma definido para el punto 1.")
    for parrafo in DEFINICION_POBLACION[2:]:
        texto_docx(documento, parrafo)
    documento.add_picture(str(CROMOSOMA_NAVE), width=Inches(6.65))
    documento.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo_figura_docx(documento, "Figura 2. Cromosoma definido para el punto 2.")

    documento.add_heading("5. Estrategia de cruce y mutación", level=1)
    for parrafo in CRUCE_MUTACION:
        texto_docx(documento, parrafo)

    documento.add_heading("6. Aplicación construida", level=1)
    texto_docx(documento, APLICACION)
    documento.add_picture(str(CAPTURA_LISTAS), width=Inches(6.65))
    documento.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo_figura_docx(documento, "Figura 3. Aplicación construida para el punto 1.")
    texto_docx(documento, APLICACION_NAVE)
    documento.add_picture(str(CAPTURA_NAVE), width=Inches(6.65))
    documento.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo_figura_docx(documento, "Figura 4. Aplicación construida para el punto 2.")
    texto_docx(
        documento,
        "En la siguiente URL se puede acceder al código de la aplicación: "
        + REPOSITORIO,
    )

    documento.add_heading("7. Conclusiones", level=1)
    for parrafo in CONCLUSIONES:
        texto_docx(documento, parrafo)

    documento.add_heading("Bibliografía", level=1)
    for referencia in BIBLIOGRAFIA:
        texto_docx(documento, referencia)

    documento.save(SALIDA_DOCX)


class DiagramaCadena(Flowable):
    def __init__(self):
        super().__init__()
        self.width = 17 * cm
        self.height = 3.5 * cm

    def draw(self):
        canvas = self.canv
        operaciones = ["+1", "×2", "×2", "+1", "×2", "+1"]
        valores = ["0 → 1", "1 → 2", "2 → 4", "4 → 5", "5 → 10", "10 → 11"]
        ancho, espacio = 1.75 * cm, 0.67 * cm
        x, y = 0.3 * cm, 1.15 * cm
        canvas.setFont("Arial-Bold", 8)
        canvas.setFillColor(colors.HexColor("#0369A1"))
        canvas.drawString(x, 2.85 * cm, "cabeza")

        for operacion, valor in zip(operaciones, valores):
            canvas.setFillColor(colors.HexColor("#E0F2FE"))
            canvas.setStrokeColor(colors.HexColor("#0284C7"))
            canvas.roundRect(x, y, ancho, 1.15 * cm, 6, fill=1, stroke=1)
            canvas.setFillColor(colors.HexColor("#0F172A"))
            canvas.setFont("Arial-Bold", 13)
            canvas.drawCentredString(x + ancho / 2, y + 0.69 * cm, operacion)
            canvas.setFont("Arial", 7)
            canvas.drawCentredString(x + ancho / 2, y + 0.25 * cm, valor)
            inicio, fin = x + ancho, x + ancho + espacio - 4
            centro = y + 0.56 * cm
            canvas.setStrokeColor(colors.HexColor("#64748B"))
            canvas.line(inicio, centro, fin, centro)
            canvas.line(fin, centro, fin - 5, centro + 3)
            canvas.line(fin, centro, fin - 5, centro - 3)
            x += ancho + espacio

        canvas.setFillColor(colors.HexColor("#CBD5E1"))
        canvas.roundRect(x - 3, y + 0.18 * cm, 1.05 * cm, 0.75 * cm, 5, fill=1)
        canvas.setFillColor(colors.HexColor("#0F172A"))
        canvas.setFont("Consolas", 8)
        canvas.drawCentredString(x + 0.47 * cm, y + 0.45 * cm, "null")
        canvas.setFillColor(colors.HexColor("#047857"))
        canvas.setFont("Arial-Bold", 8.5)
        canvas.drawString(0.3 * cm, 0.42 * cm, "Resultado = 11    Fitness = 100")


def registrar_fuentes() -> None:
    for nombre, archivo in (
        ("Arial", "arial.ttf"),
        ("Arial-Bold", "arialbd.ttf"),
        ("Arial-Italic", "ariali.ttf"),
        ("Consolas", "consola.ttf"),
    ):
        pdfmetrics.registerFont(
            TTFont(nombre, f"C:/Windows/Fonts/{archivo}")
        )


def estilos_pdf():
    base = getSampleStyleSheet()
    return {
        "titulo": ParagraphStyle(
            "Titulo",
            parent=base["Title"],
            fontName="Arial-Bold",
            fontSize=17,
            leading=21,
            alignment=TA_CENTER,
            spaceAfter=7,
        ),
        "autor": ParagraphStyle(
            "Autor",
            parent=base["Normal"],
            fontName="Arial",
            fontSize=9,
            leading=12,
            alignment=TA_CENTER,
            spaceAfter=2,
        ),
        "h1": ParagraphStyle(
            "H1",
            parent=base["Heading1"],
            fontName="Arial-Bold",
            fontSize=13,
            leading=16,
            spaceBefore=8,
            spaceAfter=5,
        ),
        "cuerpo": ParagraphStyle(
            "Cuerpo",
            parent=base["BodyText"],
            fontName="Arial",
            fontSize=9.3,
            leading=12.7,
            alignment=TA_JUSTIFY,
            spaceAfter=6,
        ),
        "ecuacion": ParagraphStyle(
            "Ecuacion",
            parent=base["Normal"],
            fontName="Arial",
            fontSize=12,
            leading=16,
            alignment=TA_CENTER,
            spaceBefore=6,
            spaceAfter=3,
        ),
        "caption": ParagraphStyle(
            "Caption",
            parent=base["Normal"],
            fontName="Arial-Italic",
            fontSize=8.5,
            leading=11,
            alignment=TA_CENTER,
            spaceAfter=7,
        ),
        "url": ParagraphStyle(
            "Url",
            parent=base["BodyText"],
            fontName="Arial",
            fontSize=9.2,
            leading=12,
            textColor=colors.HexColor("#0369A1"),
            spaceAfter=6,
        ),
    }


def numero_pagina(canvas, documento):
    canvas.saveState()
    canvas.setFont("Arial", 8)
    canvas.setFillColor(colors.HexColor("#64748B"))
    canvas.drawRightString(19.5 * cm, 1.0 * cm, str(documento.page))
    canvas.restoreState()


def imagen_pdf(ruta: Path, ancho: float) -> Image:
    with PilImage.open(ruta) as imagen:
        proporcion = imagen.height / imagen.width
    return Image(str(ruta), width=ancho, height=ancho * proporcion)


def generar_pdf() -> None:
    registrar_fuentes()
    estilos = estilos_pdf()
    documento = SimpleDocTemplate(
        str(SALIDA_PDF),
        pagesize=letter,
        rightMargin=2.0 * cm,
        leftMargin=2.0 * cm,
        topMargin=1.6 * cm,
        bottomMargin=1.5 * cm,
        title=TITULO,
        author=AUTORES,
    )
    contenido = [
        Paragraph(TITULO, estilos["titulo"]),
        Paragraph(AUTORES, estilos["autor"]),
        Paragraph(INSTITUCION, estilos["autor"]),
        Spacer(1, 0.18 * cm),
        Paragraph("1. Descripción general", estilos["h1"]),
    ]
    contenido.extend(Paragraph(p, estilos["cuerpo"]) for p in DESCRIPCION_GENERAL)
    contenido.append(Paragraph("2. Trabajos relacionados", estilos["h1"]))
    contenido.extend(Paragraph(p, estilos["cuerpo"]) for p in TRABAJOS_RELACIONADOS)

    contenido.extend([
        PageBreak(),
        Paragraph("3. Función objetivo", estilos["h1"]),
        Paragraph(FUNCION_OBJETIVO, estilos["cuerpo"]),
        Paragraph("f(c) = max(1, 100 − |O − R(c)|)", estilos["ecuacion"]),
        Paragraph("Ecuación 1. Función objetivo del punto 1.", estilos["caption"]),
        Paragraph(FUNCION_OBJETIVO_NAVE, estilos["cuerpo"]),
        Paragraph(
            "F = 4000/(1 + 0,15|x − c|) + 1500/(1 + |vy|) + "
            "600/(1 + |vx|) + 400/(1 + 2|a|) + B",
            estilos["ecuacion"],
        ),
        Paragraph("Ecuación 2. Función objetivo del punto 2.", estilos["caption"]),
        Paragraph("4. Definición de población", estilos["h1"]),
    ])
    contenido.extend(
        Paragraph(p, estilos["cuerpo"]) for p in DEFINICION_POBLACION[:2]
    )
    contenido.extend([
        DiagramaCadena(),
        Paragraph(
            "Figura 1. Cromosoma definido para el punto 1.", estilos["caption"]
        ),
    ])
    contenido.extend(
        Paragraph(p, estilos["cuerpo"]) for p in DEFINICION_POBLACION[2:]
    )
    cromosoma_nave = imagen_pdf(CROMOSOMA_NAVE, 15.8 * cm)
    contenido.extend([
        KeepTogether([
            cromosoma_nave,
            Paragraph(
                "Figura 2. Cromosoma definido para el punto 2.", estilos["caption"]
            ),
        ]),
        PageBreak(),
        Paragraph("5. Estrategia de cruce y mutación", estilos["h1"]),
    ])
    contenido.extend(Paragraph(p, estilos["cuerpo"]) for p in CRUCE_MUTACION)

    captura_listas = imagen_pdf(CAPTURA_LISTAS, 15.7 * cm)
    captura_nave = imagen_pdf(CAPTURA_NAVE, 15.7 * cm)
    contenido.extend([
        PageBreak(),
        Paragraph("6. Aplicación construida", estilos["h1"]),
        Paragraph(APLICACION, estilos["cuerpo"]),
        Spacer(1, 0.15 * cm),
        KeepTogether([
            captura_listas,
            Paragraph(
                "Figura 3. Aplicación construida para el punto 1.", estilos["caption"]
            ),
        ]),
        PageBreak(),
        Paragraph(APLICACION_NAVE, estilos["cuerpo"]),
        KeepTogether([
            captura_nave,
            Paragraph(
                "Figura 4. Aplicación construida para el punto 2.", estilos["caption"]
            ),
        ]),
        Paragraph(
            "En la siguiente URL se puede acceder al código de la aplicación: "
            f"<link href='{REPOSITORIO}'>{REPOSITORIO}</link>",
            estilos["url"],
        ),
        PageBreak(),
        Paragraph("7. Conclusiones", estilos["h1"]),
    ])
    contenido.extend(Paragraph(p, estilos["cuerpo"]) for p in CONCLUSIONES)
    contenido.append(Paragraph("Bibliografía", estilos["h1"]))
    contenido.extend(Paragraph(p, estilos["cuerpo"]) for p in BIBLIOGRAFIA)

    documento.build(
        contenido, onFirstPage=numero_pagina, onLaterPages=numero_pagina
    )


if __name__ == "__main__":
    if not CAPTURA_LISTAS.exists() or not CAPTURA_NAVE.exists():
        raise FileNotFoundError(
            "Primero genere las capturas de las dos aplicaciones"
        )
    generar_figura_cromosoma()
    generar_figura_cromosoma_nave()
    generar_docx()
    generar_pdf()
    print(f"DOCX: {SALIDA_DOCX}")
    print(f"PDF:  {SALIDA_PDF}")
